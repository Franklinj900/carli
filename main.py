import smtplib
import re
import ast
import operator
import random
import os
import json
from datetime import datetime
from io import BytesIO
from pathlib import Path
from email.mime.text import MIMEText
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from fastapi import FastAPI, Form, Request, status
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware

from database import initialize_database, load_state, save_state

# Inicialización de la aplicación
app = FastAPI()
app.add_middleware(
    SessionMiddleware,
    secret_key=os.getenv("SESSION_SECRET", "CHANGE_ME_SESSION_SECRET"),
    same_site="lax",
    https_only=os.getenv("SESSION_HTTPS_ONLY", "0") == "1",
)

BASE_DIR = Path(__file__).resolve().parent
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=BASE_DIR / "templates")

initialize_database()
players, torneos_activos, tournament_registrations, brackets_torneos = load_state()

SIMULATED_NICKNAMES = {"toti900", "nani212", "prethorian1", "car", "qkwj", "juang", "after"}
players = [p for p in players if p.get("nickname", "").strip().lower() not in SIMULATED_NICKNAMES]

MAX_TORNEO_PARTICIPANTES = 128
ADMIN_USER = os.getenv("ADMIN_USER", "admin")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "change-me")

MATRICES_FILE = BASE_DIR / "matrices_dinamicas.json"

def cargar_matrices_dinamicas():
    if MATRICES_FILE.exists():
        try:
            with open(MATRICES_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def guardar_matrices_dinamicas(matrices):
    with open(MATRICES_FILE, "w", encoding="utf-8") as f:
        json.dump(matrices, f, ensure_ascii=False, indent=4)

def limpiar_y_extraer_variables(prompt_instrucciones: str) -> List[dict]:
    lineas = prompt_instrucciones.split("\n")
    variables = []
    stopwords = ["hola", "gemini", "géminis", "necesito", "ayuda", "por favor", "gracias", "analista", "tactico", "sport", "como", "puedes", "genera", "crea", "aquí", "tienes", "archivo", "adjunto", "quiero", "ayudame", "haz"]
    
    for linea in lineas:
        sublineas = linea.split(",")
        for sub in sublineas:
            sub_clean = sub.strip()
            if not sub_clean:
                continue
            sub_lower = sub_clean.lower()
            is_conversational = any(sw in sub_lower for sw in stopwords) or sub_clean.endswith("?") or len(sub_clean.split()) > 6
            if not is_conversational:
                nombre_var = re.sub(r'[^a-zA-Z0-9áéíóúÑñ ]', '', sub_clean).strip()
                if nombre_var and len(nombre_var) > 1:
                    variables.append({"nombre": nombre_var})
    
    if not variables:
        variables = [{"nombre": "Parámetro Táctico 1"}, {"nombre": "Parámetro Táctico 2"}]
    return variables

def sincronizar_torneos_jugadores():
    for t_name, inscritos in tournament_registrations.items():
        for ins in inscritos:
            nick = ins.get("nickname", "").strip().lower()
            player = next((p for p in players if p.get("nickname", "").strip().lower() == nick), None)
            if player:
                if "torneos_participados" not in player or not isinstance(player["torneos_participados"], list):
                    player["torneos_participados"] = []
                if t_name not in player["torneos_participados"]:
                    player["torneos_participados"].append(t_name)

    for t_name, bracket_data in brackets_torneos.items():
        campeon = bracket_data.get("campeon")
        if campeon and campeon not in ("---", "Pendiente"):
            player = next((p for p in players if p.get("nickname", "").strip().lower() == campeon.strip().lower()), None)
            if player:
                if "torneos_participados" not in player or not isinstance(player["torneos_participados"], list):
                    player["torneos_participados"] = []
                if t_name not in player["torneos_participados"]:
                    player["torneos_participados"].append(t_name)
                
                if "torneos_ganados" not in player or not isinstance(player["torneos_ganados"], list):
                    player["torneos_ganados"] = []
                if t_name not in player["torneos_ganados"]:
                    player["torneos_ganados"].append(t_name)

sincronizar_torneos_jugadores()

def persistir_estado():
    sincronizar_torneos_jugadores()
    save_state(players, torneos_activos, tournament_registrations, brackets_torneos)

def usuario_de_sesion(request):
    player_id = request.session.get("player_id")
    return next((player for player in players if player.get("id") == player_id), None)

def admin_autenticado(request):
    return request.session.get("admin_authenticated") is True

def siguiente_potencia_de_dos(total):
    size = 2
    while size < total:
        size *= 2
    return size

def etiqueta_ronda(numero, total_rondas):
    restantes = total_rondas - numero + 1
    if restantes == 1:
        return "Final"
    if restantes == 2:
        return "Semifinales"
    if restantes == 3:
        return "Cuartos de final"
    return f"Ronda {numero}"

def asignar_ganador_siguiente(combates, match):
    siguiente_id = match.get("siguiente_id")
    if not siguiente_id:
        return
    siguiente = next((item for item in combates if item["id"] == siguiente_id), None)
    if not siguiente:
        return
    if match["indice"] % 2 == 0:
        siguiente["p1"] = match["ganador"]
        siguiente["seed1"] = match["seed1"]
    else:
        siguiente["p2"] = match["ganador"]
        siguiente["seed2"] = match["seed2"]

def aplicar_pases_automaticos(combates):
    cambio = True
    while cambio:
        cambio = False
        for match in combates:
            if match["ganador"] or not match.get("siguiente_id"):
                continue
            jugadores = [match["p1"], match["p2"]]
            disponibles = [jugador for jugador in jugadores if jugador not in ("---", "Pendiente")]
            if jugadores == ["---", "---"]:
                match["ganador"] = "---"
                asignar_ganador_siguiente(combates, match)
                cambio = True
            elif len(disponibles) == 1 and "---" in jugadores:
                match["ganador"] = disponibles[0]
                asignar_ganador_siguiente(combates, match)
                cambio = True

def generar_bracket(participantes, aleatorio=False):
    total = len(participantes)
    if total < 2:
        return None
    size = siguiente_potencia_de_dos(total)
    if size > MAX_TORNEO_PARTICIPANTES:
        raise ValueError(f"El torneo admite como máximo {MAX_TORNEO_PARTICIPANTES} participantes.")

    total_rondas = size.bit_length() - 1
    jugadores = [participante["nickname"] for participante in participantes]
    if aleatorio:
        random.shuffle(jugadores)
    jugadores.extend(["---"] * (size - len(jugadores)))
    combates = []
    for ronda in range(1, total_rondas + 1):
        cantidad = size // (2 ** ronda)
        for indice in range(cantidad):
            if ronda == 1:
                p1 = jugadores[indice * 2]
                p2 = jugadores[indice * 2 + 1]
                seed1 = indice * 2 + 1
                seed2 = indice * 2 + 2
            else:
                p1 = "Pendiente"
                p2 = "Pendiente"
                seed1 = "-"
                seed2 = "-"
            combates.append({
                "id": f"R{ronda}-{indice + 1}",
                "ronda": etiqueta_ronda(ronda, total_rondas),
                "ronda_numero": ronda,
                "indice": indice,
                "p1": p1,
                "seed1": seed1,
                "p2": p2,
                "seed2": seed2,
                "score1": 0,
                "score2": 0,
                "ganador": None,
                "siguiente_id": f"R{ronda + 1}-{(indice // 2) + 1}" if ronda < total_rondas else None,
            })

    aplicar_pases_automaticos(combates)
    return {
        "estado": "En Curso",
        "campeon": None,
        "combates": combates,
        "size": size,
        "total_participantes": total,
        "total_rondas": total_rondas,
        "bloqueado": False,
    }

def calcular_evaluacion(tiros_totales, tiros_arco, goles, xg, precision_pases,
                        cambio_formacion, entradas_intentadas, entradas_ganadas,
                        cursor_fallido, precision_regates, green_timing, nota_mental):
    tir = 0.0
    if tiros_totales > 0:
        base_tir = (tiros_arco / tiros_totales * 10)
        if goles > xg:
            base_tir += 1
        elif goles < xg:
            base_tir -= 1
        tir = max(0.0, min(10.0, base_tir))

    tac = 0.0
    if precision_pases > 0:
        if precision_pases > 90:
            base_tac = 9.5
        elif precision_pases >= 85:
            base_tac = 7.5
        elif precision_pases >= 80:
            base_tac = 6.0
        else:
            base_tac = 4.5
        tac = min(10.0, base_tac + (1 if cambio_formacion else 0))
    elif cambio_formacion:
        tac = 1.0

    def_stat = 0.0
    if entradas_intentadas > 0:
        base_def = (entradas_ganadas / entradas_intentadas * 10)
        def_stat = max(0.0, base_def - (1 if cursor_fallido else 0))

    mec = 0.0
    if precision_regates > 0:
        if precision_regates > 65 and green_timing:
            mec = 9.5
        elif precision_regates >= 50:
            mec = 7.5
        else:
            mec = 4.5
    elif green_timing:
        mec = 1.0

    men = max(0.0, min(10.0, nota_mental)) if nota_mental > 0 else 0.0

    componentes = [v for v in [tir, tac, def_stat, mec, men] if v > 0]
    grl = (sum(componentes) / len(componentes)) if componentes else 0.0

    return {
        "juego_tipo": "EA FC 26",
        "tir": round(tir, 1),
        "tac": round(tac, 1),
        "def": round(def_stat, 1),
        "mec": round(mec, 1),
        "men": round(men, 1),
        "grl": round(grl, 2),
    }

def calcular_evaluacion_mk(puntaje_combos, uso_defensas, flawless_victories, nota_mental):
    com = max(0.0, min(10.0, puntaje_combos)) if puntaje_combos > 0 else 0.0
    def_val = max(0.0, min(10.0, uso_defensas)) if uso_defensas > 0 else 0.0
    flaw = max(0.0, min(10.0, flawless_victories * 2.0)) if flawless_victories > 0 else 0.0
    men = max(0.0, min(10.0, nota_mental)) if nota_mental > 0 else 0.0

    componentes = [v for v in [com, def_val, flaw, men] if v > 0]
    grl = (sum(componentes) / len(componentes)) if componentes else 0.0

    return {
        "juego_tipo": "Mortal Kombat / Peleas",
        "tir": round(com, 1),
        "tac": round(def_val, 1),
        "def": round(flaw, 1),
        "mec": round(flaw, 1),
        "men": round(men, 1),
        "grl": round(grl, 2),
    }

def obtener_usuarios_registrados():
    return [
        {
            "id": p["id"],
            "nombre": p.get("name", ""),
            "apellido": p.get("surname", ""),
            "edad": p.get("age", 0),
            "nickname": p["nickname"],
            "password": p.get("password", ""),
            "telefono": p.get("phone", "+58 412-0000000"),
            "scouting": p.get("scouting", None),
            "matriz_raw": p.get("matriz_raw", None),
            "wins": p.get("wins", 0),
            "losses": p.get("losses", 0),
            "draws": p.get("draws", 0),
            "points": p.get("points", 0),
            "torneos_participados": p.get("torneos_participados", []),
            "torneos_ganados": p.get("torneos_ganados", []),
        }
        for p in players
    ]

def obtener_maximus():
    if not players:
        return None
    candidato = max(
        players, 
        key=lambda x: (x.get("scouting") or {}).get("grl", x.get("points", 0)), 
        default=None
    )
    if candidato:
        scout = candidato.get("scouting") or {}
        if scout.get("grl", 0) > 0 or candidato.get("points", 0) > 0:
            return candidato        
    return None

@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    sincronizar_torneos_jugadores()
    return templates.TemplateResponse(
        request=request,
        name="nosotros.html",
        context={
            "active_tab": "nosotros",
            "torneos_activos": torneos_activos,
            "current_user": usuario_de_sesion(request),
            "is_main": admin_autenticado(request),
            "maximus": obtener_maximus(),
        },
    )

@app.get("/nosotros", response_class=HTMLResponse)
async def nosotros_view(request: Request):
    sincronizar_torneos_jugadores()
    return templates.TemplateResponse(
        request=request,
        name="nosotros.html",
        context={
            "active_tab": "nosotros",
            "torneos_activos": torneos_activos,
            "current_user": usuario_de_sesion(request),
            "is_main": admin_autenticado(request),
            "maximus": obtener_maximus(),
        },
    )

@app.get("/torneos-activos", response_class=HTMLResponse)
async def torneos_activos_view(request: Request):
    sincronizar_torneos_jugadores()
    return templates.TemplateResponse(
        request=request,
        name="torneos.html",
        context={
            "active_tab": "torneos_activos",
            "torneos_activos": torneos_activos,
            "tournament_registrations": tournament_registrations,
            "usuarios_registrados": obtener_usuarios_registrados(),
            "brackets_torneos": brackets_torneos,
            "current_user": usuario_de_sesion(request),
            "is_main": admin_autenticado(request),
            "next_id": max([p["id"] for p in players], default=0) + 1,
            "maximus": obtener_maximus(),
        },
    )

@app.get("/perfil", response_class=HTMLResponse)
async def perfil_view(request: Request):
    sincronizar_torneos_jugadores()
    return templates.TemplateResponse(
        request=request,
        name="perfil.html",
        context={
            "active_tab": "perfil",
            "players": players,
            "torneos_activos": torneos_activos,
            "usuarios_registrados": obtener_usuarios_registrados(),
            "inscritos_torneos": tournament_registrations,
            "brackets_torneos": brackets_torneos,
            "current_user": usuario_de_sesion(request),
            "is_main": admin_autenticado(request),
            "next_id": max([p["id"] for p in players], default=0) + 1,
            "maximus": obtener_maximus(),
        },
    )

@app.get("/main", response_class=HTMLResponse)
async def main_panel(request: Request):
    sincronizar_torneos_jugadores()
    return templates.TemplateResponse(
        request=request,
        name="admin.html",
        context={
            "active_tab": "main",
            "players": players,
            "torneos_activos": torneos_activos,
            "usuarios_registrados": obtener_usuarios_registrados(),
            "inscritos_torneos": tournament_registrations,
            "brackets_torneos": brackets_torneos,
            "current_user": usuario_de_sesion(request),
            "is_main": admin_autenticado(request),
            "next_id": max([p["id"] for p in players], default=0) + 1,
            "maximus": obtener_maximus(),
            "matrices_dinamicas": cargar_matrices_dinamicas(),
        },
    )

@app.post("/main/crear-matriz-dinamica")
async def crear_matriz_dinamica(
    request: Request,
    juego_nombre: str = Form(...),
    subtitulo_abreviatura: str = Form("Criterio Táctico"),
    formula: str = Form(""),
    prompt_instrucciones: str = Form(""),
    min_val: float = Form(1.0),
    max_val: float = Form(10.0),
):
    if not admin_autenticado(request):
        return RedirectResponse(url="/main", status_code=303)
    
    variables = limpiar_y_extraer_variables(prompt_instrucciones)
    matrices = cargar_matrices_dinamicas()
    matrices[juego_nombre] = {
        "subtitulo_abreviatura": subtitulo_abreviatura,
        "formula": formula,
        "prompt_instrucciones": prompt_instrucciones,
        "variables": variables,
        "min_val": min_val,
        "max_val": max_val
    }
    guardar_matrices_dinamicas(matrices)
    return RedirectResponse(url="/main", status_code=303)

@app.get("/main/exportar-excel")
async def exportar_excel(request: Request):
    if not admin_autenticado(request):
        return RedirectResponse(url="/main", status_code=303)

    sincronizar_torneos_jugadores()
    wb = openpyxl.Workbook()

    # Estilos profesionales para el Excel
    font_titulo = Font(name="Arial", size=13, bold=True, color="FFE500")
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Arial", size=10, bold=True)
    font_normal = Font(name="Arial", size=10)

    fill_header = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    border_thin = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # 0. Pestaña de Resumen y Maximus Pretoriano
    ws_resumen = wb.active
    ws_resumen.title = "Resumen General"
    ws_resumen.views.sheetView[0].showGridLines = True

    fecha_descarga = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    maximus = obtener_maximus()

    ws_resumen.append(["COLISSEUM GAMER - REPORTE OFICIAL DE ADMINISTRACIÓN"])
    ws_resumen.cell(row=1, column=1).font = font_titulo
    ws_resumen.append([f"Fecha y Hora de Descarga del Reporte: {fecha_descarga}"])
    ws_resumen.cell(row=2, column=1).font = font_bold
    ws_resumen.append([])

    if maximus:
        scout_m = maximus.get("scouting") or {}
        ws_resumen.append(["MAXIMUS PRETORIANO ACTUAL"])
        ws_resumen.cell(row=4, column=1).font = font_header
        ws_resumen.cell(row=4, column=1).fill = fill_header

        datos_max = [
            ["Nickname", maximus.get("nickname")],
            ["Nombre Completo", f"{maximus.get('name', '')} {maximus.get('surname', '')}"],
            ["Edad", maximus.get("age", 0)],
            ["Teléfono", maximus.get("phone", "")],
            ["Puntos Totales", maximus.get("points", 0)],
            ["GRL / Calificación", scout_m.get("grl", 0.0)],
            ["Juego / Matriz", scout_m.get("juego_tipo", "EA FC 26")],
            ["Récord de Victorias", f"{maximus.get('wins', 0)}W - {maximus.get('losses', 0)}L - {maximus.get('draws', 0)}E"]
        ]
        for row_idx, item in enumerate(datos_max, start=5):
            ws_resumen.append(item)
            ws_resumen.cell(row=row_idx, column=1).font = font_bold
            ws_resumen.cell(row=row_idx, column=2).font = font_normal
    else:
        ws_resumen.append(["No hay Maximus Pretoriano definido todavía."])

    # 1. Pestaña de Gladiadores
    ws1 = wb.create_sheet(title="Gladiadores")
    ws1.views.sheetView[0].showGridLines = True
    headers_g = ["ID", "Nombre", "Apellido", "Edad", "Nickname (Usuario)", "Contraseña", "Teléfono", "Victorias", "Derrotas", "Empates", "Puntos", "Torneos Participados", "Torneos Ganados"]
    ws1.append(headers_g)
    for col_idx in range(1, len(headers_g) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, p in enumerate(players, start=2):
        row_data = [
            p.get("id"),
            p.get("name"),
            p.get("surname"),
            p.get("age", 0),
            p.get("nickname"),
            p.get("password"),
            p.get("phone"),
            p.get("wins", 0),
            p.get("losses", 0),
            p.get("draws", 0),
            p.get("points", 0),
            ", ".join(p.get("torneos_participados", [])),
            ", ".join(p.get("torneos_ganados", []))
        ]
        ws1.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            c = ws1.cell(row=idx, column=col_idx)
            c.font = font_normal
            c.border = border_thin
            if idx % 2 == 0:
                c.fill = fill_zebra

    # 2. Pestaña de Rendimiento
    ws2 = wb.create_sheet(title="Rendimiento")
    ws2.views.sheetView[0].showGridLines = True
    headers_r = ["Nickname", "Matriz / Juego", "TIR / Combos", "TAC / Defensas", "DEF / Flawless", "MEC", "MEN", "GRL", "Habilidades"]
    ws2.append(headers_r)
    for col_idx in range(1, len(headers_r) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, p in enumerate(players, start=2):
        scout = p.get("scouting")
        if scout:
            row_data = [
                p.get("nickname"),
                scout.get("juego_tipo", "EA FC 26"),
                scout.get("tir"),
                scout.get("tac"),
                scout.get("def"),
                scout.get("mec"),
                scout.get("men"),
                scout.get("grl"),
                scout.get("habilidades", "")
            ]
            ws2.append(row_data)
            for col_idx in range(1, len(row_data) + 1):
                c = ws2.cell(row=idx, column=col_idx)
                c.font = font_normal
                c.border = border_thin
                if idx % 2 == 0:
                    c.fill = fill_zebra

    # 3. Pestaña de Torneos y Participantes
    ws3 = wb.create_sheet(title="Torneos")
    ws3.views.sheetView[0].showGridLines = True
    headers_t = ["Torneo", "Subtítulo", "Tipo de Juego", "Fecha de Inicio / Publicación", "Formato", "Premio", "Máx. Participantes", "Estado", "Campeón", "Total Inscritos", "Participantes Inscritos"]
    ws3.append(headers_t)
    for col_idx in range(1, len(headers_t) + 1):
        cell = ws3.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, t in enumerate(torneos_activos, start=2):
        t_name = f"{t['titulo']} ({t['subtitulo']})" if t.get("subtitulo") else t['titulo']
        bracket = brackets_torneos.get(t_name, {})
        inscritos = tournament_registrations.get(t_name, [])
        nicks_inscritos = ", ".join([ins.get("nickname", "") for ins in inscritos])

        row_data = [
            t.get("titulo"),
            t.get("subtitulo", ""),
            t.get("juego_tipo", "EA FC 26"),
            t.get("fecha_inicio", "Por definir"),
            t.get("formato", "Eliminación Directa"),
            t.get("premio", "Por definir"),
            t.get("max_participantes", 128),
            bracket.get("estado", "Pendiente"),
            bracket.get("campeon", "Pendiente"),
            len(inscritos),
            nicks_inscritos
        ]
        ws3.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            c = ws3.cell(row=idx, column=col_idx)
            c.font = font_normal
            c.border = border_thin
            if idx % 2 == 0:
                c.fill = fill_zebra

    # Autoajustar ancho de columnas en todas las hojas
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    val_str = str(cell.value)
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    headers = {'Content-Disposition': 'attachment; filename="Colisseum_Gamer_Database_Completa.xlsx"'}
    return StreamingResponse(output, headers=headers, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.post("/registrar-gladiador")
async def registrar_gladiador(
    request: Request,
    nombre: str = Form(...),
    apellido: str = Form(...),
    edad: int = Form(...),
    nickname: str = Form(...),
    telefono: str = Form(...),
    password: str = Form(...),
):
    nickname_existe = any(
        p.get("nickname", "").strip().lower() == nickname.strip().lower()
        for p in players
    )

    referer = request.headers.get("referer", "/torneos-activos")
    base_referer = referer.split("?")[0]

    if nickname_existe:
        return RedirectResponse(url=f"{base_referer}?registro_error=true", status_code=303)

    nuevo_gladiador = {
        "id": max([p["id"] for p in players], default=0) + 1,
        "name": nombre,
        "surname": apellido,
        "age": edad,
        "nickname": nickname,
        "password": password,
        "points": 0,
        "wins": 0,
        "losses": 0,
        "draws": 0,
        "phone": telefono,
        "scouting": None,
        "matriz_raw": None,
        "torneos_participados": [],
        "torneos_ganados": [],
    }
    players.append(nuevo_gladiador)
    persistir_estado()

    return RedirectResponse(url=f"{base_referer}?registro_exito=true", status_code=303)

@app.post("/main-login")
async def main_login(
    request: Request,
    main_user: str = Form(...),
    main_pass: str = Form(...)
):
    if main_user == ADMIN_USER and main_pass == ADMIN_PASSWORD:
        request.session["admin_authenticated"] = True
        return RedirectResponse(url="/main", status_code=303)
    return RedirectResponse(url="/main?error=admin_auth", status_code=303)

@app.get("/main-logout")
async def main_logout(request: Request):
    request.session.pop("admin_authenticated", None)
    return RedirectResponse(url="/main", status_code=303)

@app.post("/main/crear-torneo")
async def crear_torneo(
    request: Request,
    titulo: str = Form(...),
    subtitulo: str = Form(""),
    juego_tipo: str = Form("ea_fc_26"),
    fecha_inicio: str = Form(""),
    formato: str = Form("Eliminación Directa"),
    premio: str = Form("Por definir"),
    max_participantes: int = Form(128),
    reglas: str = Form(""),
):
    if not admin_autenticado(request):
        return RedirectResponse(url="/main", status_code=303)

    nuevo_id = max([t.get("id", 0) for t in torneos_activos], default=0) + 1
    nuevo_torneo = {
        "id": nuevo_id,
        "titulo": titulo,
        "subtitulo": subtitulo,
        "juego_tipo": juego_tipo,
        "fecha_inicio": fecha_inicio,
        "formato": formato,
        "premio": premio,
        "max_participantes": max_participantes,
        "reglas": reglas,
    }
    torneos_activos.append(nuevo_torneo)

    nombre_key = f"{titulo} ({subtitulo})" if subtitulo else titulo
    if nombre_key not in tournament_registrations:
        tournament_registrations[nombre_key] = []

    persistir_estado()
    return RedirectResponse(url="/main", status_code=303)

@app.post("/main/eliminar-torneo")
async def eliminar_torneo(request: Request, titulo_torneo: str = Form(...)):
    if not admin_autenticado(request):
        return RedirectResponse(url="/main", status_code=303)

    global torneos_activos
    torneos_activos[:] = [
        t for t in torneos_activos
        if (f"{t['titulo']} ({t['subtitulo']})" if t.get("subtitulo") else t["titulo"]) != titulo_torneo
    ]
    
    inscritos_previos = tournament_registrations.get(titulo_torneo, [])
    for ins in inscritos_previos:
        nick = ins.get("nickname", "").strip().lower()
        player = next((p for p in players if p.get("nickname", "").strip().lower() == nick), None)
        if player:
            if "torneos_participados" not in player or not isinstance(player["torneos_participados"], list):
                player["torneos_participados"] = []
            if titulo_torneo in player["torneos_participados"]:
                player["torneos_participados"].remove(titulo_torneo)

    tournament_registrations.pop(titulo_torneo, None)
    brackets_torneos.pop(titulo_torneo, None)

    persistir_estado()
    return RedirectResponse(url="/main", status_code=303)

@app.post("/main/eliminar-usuario")
async def eliminar_usuario(request: Request, user_id: int = Form(...)):
    if not admin_autenticado(request):
        return RedirectResponse(url="/main", status_code=303)

    global players
    players[:] = [p for p in players if p.get("id") != user_id]

    if request.session.get("player_id") == user_id:
        request.session.pop("player_id", None)

    persistir_estado()
    return RedirectResponse(url="/main", status_code=303)

@app.post("/main/reorganizar-cuadrante")
async def reorganizar_cuadrante(request: Request, tournament_name: str = Form(...)):
    if not admin_autenticado(request):
        return RedirectResponse(url="/main", status_code=303)

    inscritos = tournament_registrations.get(tournament_name, [])
    if len(inscritos) < 2:
        return RedirectResponse(url="/main?cuadrante=insuficientes", status_code=303)

    try:
        brackets_torneos[tournament_name] = generar_bracket(inscritos, aleatorio=True)
        persistir_estado()
    except Exception:
        pass

    return RedirectResponse(url="/main", status_code=303)

@app.post("/login")
async def login(request: Request, nickname: str = Form(...), password: str = Form(...)):
    for player in players:
        if player["nickname"].strip().lower() == nickname.strip().lower() and password == player["password"]:
            request.session["player_id"] = player["id"]
            return RedirectResponse(url="/perfil", status_code=303)
    return RedirectResponse(url="/perfil?error=1", status_code=303)

@app.get("/logout")
async def logout(request: Request):
    request.session.pop("player_id", None)
    return RedirectResponse(url="/nosotros", status_code=303)

@app.post("/inscribir-torneo")
async def inscribir_torneo(
    request: Request,
    tournament_name: str = Form(...),
    nickname: str = Form(...),
):
    jugador_registrado = next(
        (p for p in players if p.get("nickname", "").strip().lower() == nickname.strip().lower()),
        None
    )

    referer = request.headers.get("referer", "/torneos-activos")
    base_referer = referer.split("?")[0]

    if not jugador_registrado:
        return RedirectResponse(url=f"{base_referer}?error=no_registrado", status_code=303)

    if tournament_name not in tournament_registrations:
        tournament_registrations[tournament_name] = []

    lista_inscritos = tournament_registrations[tournament_name]

    if len(lista_inscritos) >= MAX_TORNEO_PARTICIPANTES:
        return RedirectResponse(url=f"{base_referer}?error=limite", status_code=303)

    duplicado = any(
        ins["nickname"].strip().lower() == jugador_registrado["nickname"].strip().lower()
        for ins in lista_inscritos
    )

    if duplicado:
        return RedirectResponse(url=f"{base_referer}?error=duplicado", status_code=303)

    lista_inscritos.append({
        "name": jugador_registrado.get("name", ""),
        "surname": jugador_registrado.get("surname", ""),
        "nickname": jugador_registrado["nickname"],
        "phone": jugador_registrado.get("phone", "+58 412-0000000"),
    })

    if "torneos_participados" not in jugador_registrado:
        jugador_registrado["torneos_participados"] = []
    if tournament_name not in jugador_registrado["torneos_participados"]:
        jugador_registrado["torneos_participados"].append(tournament_name)

    if len(lista_inscritos) >= 2:
        try:
            brackets_torneos[tournament_name] = generar_bracket(lista_inscritos, aleatorio=True)
        except Exception:
            pass

    persistir_estado()
    return RedirectResponse(url=f"{base_referer}?exito=true", status_code=303)

@app.post("/eliminar-inscrito")
async def eliminar_inscrito(
    request: Request,
    tournament_name: str = Form(...),
    nickname: str = Form(...),
):
    if tournament_name in tournament_registrations:
        tournament_registrations[tournament_name] = [
            ins for ins in tournament_registrations[tournament_name]
            if ins["nickname"].strip().lower() != nickname.strip().lower()
        ]
        lista_inscritos = tournament_registrations[tournament_name]
        if len(lista_inscritos) >= 2:
            try:
                brackets_torneos[tournament_name] = generar_bracket(lista_inscritos, aleatorio=True)
            except Exception:
                pass
        else:
            brackets_torneos.pop(tournament_name, None)
        persistir_estado()

    if admin_autenticado(request):
        return RedirectResponse(url="/main", status_code=303)
    return RedirectResponse(url="/torneos-activos", status_code=303)

@app.post("/main/actualizar-bracket")
async def actualizar_bracket(
    request: Request,
    tournament_name: str = Form(...),
    combate_id: str = Form(...),
    score1: int = Form(...),
    score2: int = Form(...),
    ganador_opcion: str = Form(...),
):
    if not admin_autenticado(request):
        return RedirectResponse(url="/torneos-activos", status_code=303)

    if tournament_name in brackets_torneos:
        bracket_data = brackets_torneos[tournament_name]
        combates = bracket_data.get("combates", [])
        match = next((c for c in combates if c["id"] == combate_id), None)

        if match:
            match["score1"] = score1
            match["score2"] = score2
            ganador = match["p1"] if ganador_opcion == "1" else match["p2"]
            match["ganador"] = ganador
            asignar_ganador_siguiente(combates, match)
            aplicar_pases_automaticos(combates)

            if match.get("ronda") == "Final" and ganador and ganador not in ("---", "Pendiente"):
                bracket_data["campeon"] = ganador
                jugador_campeon = next((p for p in players if p.get("nickname", "").strip().lower() == ganador.strip().lower()), None)
                if jugador_campeon:
                    jugador_campeon["wins"] = jugador_campeon.get("wins", 0) + 1
                    
                    if "torneos_ganados" not in jugador_campeon:
                        jugador_campeon["torneos_ganados"] = []
                    if tournament_name not in jugador_campeon["torneos_ganados"]:
                        jugador_campeon["torneos_ganados"].append(tournament_name)

        persistir_estado()

    return RedirectResponse(url="/torneos-activos", status_code=303)

@app.post("/main/evaluar-gladiador")
async def evaluar_gladiador(
    request: Request,
    player_id: int = Form(...),
    juego_tipo: str = Form("ea_fc_26"),
    wins: int = Form(0),
    losses: int = Form(0),
    draws: int = Form(0),
    tiros_totales: int = Form(0),
    tiros_arco: int = Form(0),
    goles: int = Form(0),
    xg: float = Form(0.0),
    precision_pases: float = Form(0.0),
    cambio_formacion: str = Form(None),
    entradas_intentadas: int = Form(0),
    entradas_ganadas: int = Form(0),
    cursor_fallido: str = Form(None),
    precision_regates: float = Form(0.0),
    green_timing: str = Form(None),
    nota_mental: float = Form(0.0),
    habilidades: str = Form(""),
):
    if not admin_autenticado(request):
        return RedirectResponse(url="/main", status_code=303)

    matrices = cargar_matrices_dinamicas()
    form_data = await request.form()

    if juego_tipo == "mortal_kombat":
        evaluacion = calcular_evaluacion_mk(
            puntaje_combos=float(tiros_totales),
            uso_defensas=float(precision_pases),
            flawless_victories=float(goles),
            nota_mental=nota_mental
        )
    elif juego_tipo.startswith("custom_"):
        juego_nombre = juego_tipo.replace("custom_", "")
        matriz_info = matrices.get(juego_nombre, {"formula": "", "variables": [], "min_val": 1.0, "max_val": 10.0})
        variables = matriz_info.get("variables", [])
        if not variables:
            variables = limpiar_y_extraer_variables(matriz_info.get("prompt_instrucciones", ""))
            
        formula = matriz_info.get("formula", "").strip()
        min_val = float(matriz_info.get("min_val", 1.0))
        max_val = float(matriz_info.get("max_val", 10.0))
        
        local_vars = {}
        detalles = []
        
        for idx, v in enumerate(variables):
            val_str = form_data.get(f"custom_param_{idx}", form_data.get(f"custom_var_{idx}", "0"))
            try:
                val = float(val_str)
            except ValueError:
                val = 0.0
            var_nombre_key = re.sub(r'[^a-zA-Z0-9_]', '_', v.get("nombre", f"var_{idx}"))
            if not var_nombre_key or var_nombre_key[0].isdigit():
                var_nombre_key = f"v_{idx}"
            local_vars[var_nombre_key] = val
            local_vars[f"var_{idx}"] = val
            detalles.append({"nombre": v.get("nombre"), "valor": val, "key": var_nombre_key})
        
        resultado_calculado = 0.0
        if formula:
            try:
                import math
                safe_globals = {"__builtins__": {}, "math": math, "abs": abs, "max": max, "min": min, "round": round}
                resultado_calculado = float(eval(formula, safe_globals, local_vars))
            except Exception:
                resultado_calculado = sum(local_vars.values())
        else:
            resultado_calculado = sum(local_vars.values()) / len(variables) if variables else nota_mental
        
        grl_val = max(min_val, min(max_val, resultado_calculado))
        
        if max_val > 10:
            grl_normalizado = 1.0 + (grl_val / max_val) * 9.0
        else:
            grl_normalizado = grl_val

        evaluacion = {
            "juego_tipo": juego_nombre,
            "tir": round(grl_normalizado, 1),
            "tac": round(grl_normalizado, 1),
            "def": round(grl_normalizado, 1),
            "mec": round(grl_normalizado, 1),
            "men": round(nota_mental, 1),
            "grl": round(grl_val, 2),
            "escala_max": max_val,
            "detalles": detalles
        }
    else:
        evaluacion = calcular_evaluacion(
            tiros_totales=tiros_totales,
            tiros_arco=tiros_arco,
            goles=goles,
            xg=xg,
            precision_pases=precision_pases,
            cambio_formacion=bool(cambio_formacion),
            entradas_intentadas=entradas_intentadas,
            entradas_ganadas=entradas_ganadas,
            cursor_fallido=bool(cursor_fallido),
            precision_regates=precision_regates,
            green_timing=bool(green_timing),
            nota_mental=nota_mental,
        )
    
    evaluacion["habilidades"] = habilidades

    matriz_raw = {
        "wins": wins,
        "losses": losses,
        "draws": draws,
        "tiros_totales": tiros_totales,
        "tiros_arco": tiros_arco,
        "goles": goles,
        "xg": xg,
        "precision_pases": precision_pases,
        "cambio_formacion": bool(cambio_formacion),
        "entradas_intentadas": entradas_intentadas,
        "entradas_ganadas": entradas_ganadas,
        "cursor_fallido": bool(cursor_fallido),
        "precision_regates": precision_regates,
        "green_timing": bool(green_timing),
        "nota_mental": nota_mental,
        "habilidades": habilidades,
        "juego_tipo": juego_tipo
    }

    for player in players:
        if player.get("id") == player_id:
            player["wins"] = wins
            player["losses"] = losses
            player["draws"] = draws
            player["scouting"] = evaluacion
            player["matriz_raw"] = matriz_raw
            player["points"] = int(evaluacion["grl"] * 100)
            break

    persistir_estado()
    return RedirectResponse(url="/main", status_code=status.HTTP_303_SEE_OTHER)

@app.post("/enviar-mensaje")
async def enviar_mensaje(
    request: Request,
    nombre: str = Form(...),
    email: str = Form(...),
    asunto: str = Form(...),
    mensaje: str = Form(...),
):
    return RedirectResponse(url="/nosotros?mensaje_enviado=true", status_code=303)